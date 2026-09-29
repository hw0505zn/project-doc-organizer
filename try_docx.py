import docx
d = docx.Document("data/raw/省建院周报2025.7.20（第5期）.docx")
print("\n".join(p.text for p in d.paragraphs))
# 若含表格，段落取不到，要单独取：
for t in d.tables:
    for row in t.rows:
        print([c.text for c in row.cells])